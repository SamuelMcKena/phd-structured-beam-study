#!/usr/bin/env python3
"""Academic DOCX source with native editable equations and evidence-derived values."""
import argparse,json,subprocess
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
BREAK='\n\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n\n'


def run(data,figs,out):
    out.mkdir(parents=True,exist_ok=True)
    sweep=pd.read_csv(data/'sweep.csv');fine=pd.read_csv(data/'post_fine_metrics.csv')
    roles=pd.read_csv(data/'roles.csv');conv=pd.read_csv(data/'convergence.csv')
    unit=pd.read_csv(data/'unit_cell.csv');top=pd.read_csv(data/'topology_and_ring.csv')
    zones=pd.read_csv(data/'axial_zones.csv');checks=pd.read_csv(data/'propagation_checks.csv')
    fit=json.loads((data/'scaling_fit.json').read_text());manifest=json.loads((data/'manifest.json').read_text())
    stress=pd.read_csv(data/'extreme_stress.csv')
    def row(L,w=250,d=.5,axis='x',df=sweep):
        return df[(df.L==L)&(df.w_px==w)&(df.offset_px==d)&(df.axis==axis)].iloc[0]
    b10=row(10);b20=row(20);f10=row(10,df=fine);f20=row(20,df=fine);s20=row(20,w=50,df=fine)
    pages=[];fignum=0
    def page(title,text):pages.append('# '+title+'\n\n'+text.strip()+'\n')
    def figure(name,caption,width=6.55):
        nonlocal fignum
        fignum+=1
        return f'\n\n![]({(figs/(name+".png")).resolve()}){{width={width}in}}\n\nFigure {fignum}. {caption}\n\n'
    def table(headers,rows):
        return '\n|'+'|'.join(headers)+'|\n|'+'|'.join(['---']*len(headers))+'|\n'+''.join('|'+'|'.join(map(str,r))+'|\n' for r in rows)+'\n'
    def pct(x):return f'{100*x:.2f}%'
    page('Subpixel registration in a dual SLM vortex Bessel system',rf'''
Samuel McKenna  
Applied Optics and Photonics Group, Heriot-Watt University  
Technical research report, October 2026

## Abstract

Subpixel registration is the position of a mutually centred optical beam and continuous hologram relative to a physical SLM pixel lattice. It is distinct from beam–hologram decentre. This report audits the existing registration implementation and evaluates a phase-only dual-SLM reference with an 8 μm pitch and a 20-pixel blaze on **both** panels. The primary study adjusts the ideal Fourier-plane iris to each charge and beam size, then holds it fixed throughout that registration sweep. A separate fixed-iris control distinguishes registration from rejection of the intended vortex spectrum.

The historical explicit propagation route fails the repository's numerical power-loss gate: individual nominal free-space steps discard approximately 9–55% of power in the audited cases. Its previous quantitative propagated conclusions are therefore withdrawn. The replacement reference integrates the piecewise constant panel commands over the actual intersections of the two pixel grids and performs ideal-4F selection in source spatial-frequency coordinates. Independently resolved raster tests, illumination quadrature, increased cone sampling and a propagation-window control support the new calculation.

For the approximately 2 mm beam radius, $w=250p$, half-pixel x registration of the vortex-owning panel gives selected-field infidelities of {b10.raw_infidelity:.2e} and {b20.raw_infidelity:.2e} for $L=10$ and 20. In the conditional nominal cone model, the corresponding propagated peak changes are {pct(f10.peak_ratio-1)} and {pct(f20.peak_ratio-1)}, while power-normalised intensity residuals are {pct(f10.power_normalised_residual_l2)} and {pct(f20.power_normalised_residual_l2)}. The rings remain robust. At $L=20$, reducing the radius to 50 pixels increases the propagated shape residual to {pct(s20.power_normalised_residual_l2)}; at 25 pixels it reaches {pct(stress.iloc[-1].power_normalised_residual_l2)}.

With flat correction, identical panel responses and no measured inter-SLM transfer, exchanging the vortex owner is exactly a panel relabelling. The model cannot rank the two physical architectures. The missing measurements are the inter-panel transfer and coordinate mapping, the individual lens and iris geometry, the axicon cone, and panel-specific phase response and correction maps. The local registration result is reproducible; absolute laboratory beam dimensions and an architecture winner remain calibration-dependent.
''')
    page('1 Motivation and definition of registration',r'''
The practical question is whether translating the illuminated vortex structure by a fraction of an SLM pixel changes the optical field delivered to the axicon and the propagated ring system. A useful answer must identify which change arises locally at the sampled phase mask, which survives order selection, and which appears in a camera intensity image. Complex-field change and visible intensity change are not interchangeable.

At each panel let $b$ denote the beam centre, $h$ the intended continuous hologram centre and $g$ a pixel-boundary lattice origin. The two geometric quantities are

$$d_{\mathrm{decentre}}=b-h,\qquad \boldsymbol{\rho}=\frac{h-g}{p}\pmod{1}.$$

In a pure registration sweep, $b=h=0$ remains fixed. Setting panel translation to $-\boldsymbol{\rho}p$ and electronic pattern offset to $+\boldsymbol{\rho}p$ keeps the continuous hologram centre fixed while the pixel boundaries move underneath it. Moving only the electronic vortex origin would instead introduce decentre. Figure 1 makes this distinction explicit.
'''+figure('01_registration_geometry',r'The beam and intended hologram centre remain coincident at the red marker. Grey physical pixel boundaries move by $-\boldsymbol{\rho}p$. The blue region is the illuminated central part of a broad beam; its outer footprint lies outside the magnified view. In this boundary-origin convention, zero registration is a corner, a half-pixel shift along one coordinate is an edge, and $(1/2,1/2)$ is a pixel centre. The five x positions correspond to 0, 1, 2, 3 and 4 μm.')+r'''
The absolute reference registration is a convention, not a measured laboratory state. Differences are always relative to the explicitly declared zero-registration pixelated field. Whole-pitch periodicity describes the interior lattice; finite panel-edge motion is not the mechanism being studied.
''')
    page('2 Experimental architecture and known parameters',r'''
The phase-allocation comparison places the vortex on SLM1 in Architecture A and on SLM2 in Architecture B. Both panels retain a blaze in both architectures. Figure 2 distinguishes this allocation from the unknown physical transfer between the panels. Correction is exactly zero in the primary run: this is the **FLAT-CORRECTION PHYSICS-ISOLATION BASELINE**.
'''+figure('02_architectures','Architecture A uses SLM1 for vortex plus blaze and SLM2 for correction plus blaze. Architecture B swaps vortex and correction ownership. The inter-SLM transfer is unknown; the local reference uses identity transfer as a declared mathematical baseline, not a measured imaging relation.',6.45)+table(['Quantity','Value or status'],[
        ['Wavelength','1029 nm model value; approximately 1030 nm bench'],
        ['SLM pitch and phase addressing','8 μm; ideal 256-level phase response'],
        ['Blaze on each panel','6250 cycles/m; 20 pixels per period'],
        ['Bench beam radius','Approximately 2 mm = 250 pixels'],
        ['Lens separation','User-reported 300 mm; individual focal lengths unknown'],
        ['Physical axicon','User-described 20° Thorlabs optic; exact identity and orientation unverified']])+r'''
HOLOEYE specifies the PLUTO-2.1 lattice dimensions and pitch [1]. A 300 mm separation fixes $f_1+f_2$ only for an ideal 4F arrangement; it does not establish $f_1=f_2=150$ mm or unity magnification. Those equal-f values are illustrative coordinate conversions only. The primary filter is defined directly in source frequency units, so its local result does not require an invented focal length.
''')
    page('3 Code and physics audit',r'''
The audit covered the registration modules, metrics, architecture runner, plotting and report tools, their focused tests, repository hardware and evidence registries, the uploaded publication-study source, the archived v0.6 GUI, the practical-system volume, and the calibration and q20 retrieval bundles. Historical evidence was retained. The repositories contain nominal and demonstration geometry, but no validated inter-SLM distance, relay prescription, parity, rotation or magnification for this experiment.

The compensated shift itself is correct. Both blazes are present in the current architecture source. The previous one-blaze interpretation was invalid because it compared different physical phase commands rather than vortex ownership under matched hardware. Correcting that error restores role-exchange symmetry; it does not justify an architecture winner.

The remaining propagation problem is independent of that correction. At $N_{\mathrm{fine}}=2500$ over 10 mm and $N_{\mathrm{relay}}=768$, the audited band-limited ASM steps discard power in nominal free space. The losses below occur before or after the physical iris and must not be counted as optical transmission.
'''+table(['Audit case','First free-space step loss','Other large step loss'],[
        ['L = 20, w = 50 px, half-pixel x','Approximately 55%','Approximately 3%'],
        ['L = 20, w = 250 px, half-pixel x','Approximately 9%','Approximately 22% and 41%']])+r'''
These losses violate the repository's 5% numerical drift limit. The old 0.267 stress-case infidelity and its propagated peak-collapse narrative are not adopted as quantitative results. The legacy implementation now records each propagation loss separately and exposes a claim-eligibility flag. It remains available for forensic reproduction.

An iris at a Fourier plane truncates the **source angular spectrum**. It does not generally band-limit the spatial spectrum of the field expressed on that physical Fourier plane. Coarsening there therefore requires a separate sampling argument. The downstream lens phase must also be resolved. The independent reference removes these sampled lens/plane operations from the ideal-4F baseline and validates source-frequency selection directly.
''')
    page('4 Numerical model and the adjusted iris',r'''
The input scalar field is $A(x,y)=\exp[-(x^2+y^2)/w^2]$, where $w$ is the $1/e$ amplitude radius and $1/e^2$ intensity radius. A centre-sampled phase command is consistent with the inspected GUI's direct coordinate evaluation. The circular area-average command in the historical runner is a different hologram-writing convention; it is not a measured LC response and is not used as the primary command.

For Architecture A,

$$\phi_1=L\operatorname{atan2}(y,x)+2\pi c x,\qquad\phi_2=2\pi c x.$$

Each phase is evaluated at the centre of its own physical pixel and quantised to 256 uniformly spaced phase values. At a point in the overlap of the two grids,

$$E_{12}(x,y)=\mathrm{FF}\,A(x,y)\exp[i\phi_{1,\mathrm{pix}}(x,y)+i\phi_{2,\mathrm{pix}}(x,y)].$$

Here each panel contributes amplitude $\sqrt{\mathrm{FF}}$, with $\mathrm{FF}=0.93$. This throughput-only approximation does not simulate dead-space leakage, fringing, a measured LUT or an unmodulated zero order. No additional diffraction-efficiency factor is applied.

The union of the two lattices partitions the plane into rectangles on which the phase product is constant. Gaussian illumination is approximated locally and refined by subdividing those rectangles. A constant-amplitude rectangle centred on $(x_a,y_a)$ has the analytic Fourier integral

$$\widetilde E_a(f_x,f_y)=E_a\Delta x_a\Delta y_a\,\operatorname{sinc}(f_x\Delta x_a)\operatorname{sinc}(f_y\Delta y_a)e^{-i2\pi(f_xx_a+f_yy_a)}.$$

The convention is $\operatorname{sinc}(u)=\sin(\pi u)/(\pi u)$. Summing these integrals preserves fractional lattice positions without rounding them onto a fine computational raster. Fourier lattice sums are evaluated by FFT; explicit sinc and origin factors retain the true physical frequencies, including additional replica windows when needed.

The two identical carrier ramps give an effective-channel centre $(2c,0)$. An ideal circular passband of radius $B$ is applied there and the carrier is removed in an unfolded image frame. The primary adjusted-iris rule is

$$B(L,w)=\max\!\left(2500,\frac{5L}{2\pi w}\right)\ \mathrm{cycles/m}.$$

This is a diagnostic design choice. The iris changes between charge/size cases but is held fixed across every offset, direction and architecture within a case. For $L=20$, $B$ is 39.8 kcycles/m at 50 pixels and 7.96 kcycles/m at 250 pixels. A widely opened physical iris may admit zero or unwanted orders absent from this selected-channel model. The calculation therefore isolates registration; it does not assert laboratory order purity with those apertures.
''')
    page('5 Validation and convergence',r'''
The inherited 41 focused tests passed before modifications. The extended suite contains 53 passing tests, including an independent directly resolved raster comparison at 4, 8 and 16 samples per physical pixel. The production model does not rely on those raster points to locate the boundaries: its rectangular integrals represent the 8 μm lattice directly.

The gates check zero-state consistency, whole-pitch periodicity, invariant beam–hologram centring, both 20-pixel blazes, the selected-order frequency window, no-carrier rotational symmetry, permitted carrier anisotropy, exact A/B role correspondence, quadrature convergence, Parseval consistency and phase-only axicon intensity invariance. Figure 3 compares the illumination quadrature refinements. Headline fields use two subdivisions per intersection interval; four subdivisions provide the independent convergence control.
'''+figure('07_convergence','Half-pixel x registration at L = 10 and 20, for 50- and 250-pixel beam radii. The commanded pixel phases stay fixed while illumination integration is refined. Infidelity and transmitted power are separate convergence quantities.')+rf'''
The downstream cone field is propagated with scalar exact angular-spectrum propagation on a grid resolving the conical radial period. The nominal 20° base-angle model has period {manifest['cone_radial_period_m']*1e6:.2f} μm. Headline XY fields use approximately 1.25 μm sampling; the coarser approximately 2 μm run supplies the axial scan. The largest recorded relative propagation power drift in the primary scan is {checks.propagation_power_drift_fraction.max():.2e}. Conservation alone does not prove spatial accuracy: a larger-window control and increased cone sampling are also retained. Absolute peak values remain sampling-dependent, while the paired registration differences are substantially more stable.
''')
    page('6 Local charge and beam size dependence',r'''
Figure 4 separates the local two-panel field from the field passed by the adjusted iris. The local masks are phase-only: apart from the declared mean throughput and panel aperture, their immediate intensity envelope does not change with registration. The changed quantity is the phase sampled by each pixel, especially near the vortex singularity and the blaze steps. Intensity differences arise after propagation or Fourier filtering.
'''+figure('03_charge_and_beam_size',r'Half-pixel x registration of the vortex-plus-blaze panel. Left: raw complex-field infidelity immediately after both panels, integrated with the Gaussian intensity weight. Right: infidelity after ideal 4F selection, with the iris adjusted by the stated rule. Small radii are deliberate stress cases; 250 pixels is the bench-size case.')+rf'''
At 250 pixels, local half-pixel x infidelity is {pct(b10.local_infidelity)} for $L=10$ and {pct(b20.local_infidelity)} for $L=20$. This should not be interpreted as several-percent final intensity distortion. The staircase carrier contributes a nonzero floor even at $L=0$, and much of the pixel-scale phase change is rejected by selection. The selected-field infidelities fall to {b10.raw_infidelity:.2e} and {b20.raw_infidelity:.2e}.

Shrinking the beam increases the fraction of illuminated power that encounters poorly resolved azimuthal phase near the origin. Increasing charge increases its local gradient, $|\nabla\phi_{{\mathrm{{vortex}}}}|=|L|/r$. The same physical pitch therefore becomes more restrictive. Individual filtered curves need not increase monotonically with charge because the chosen passband, square-lattice harmonics and the zero-registration reference all matter. A single charge-only tolerance is not supported.
''')
    page('6.3 Direction and carrier anisotropy',r'''
The blaze is x-directed in the declared effective-channel coordinates. Registration along x changes the location of its staircase discontinuities; a y shift leaves a blaze-only command unchanged in the infinite interior lattice. A vortex samples both coordinates. Directional registration is consequently a physical role-dependent effect, not an isotropic displacement magnitude.
'''+figure('04_direction_and_power',r'L = 20 and half-pixel registration. The diagonal means $(\Delta x,\Delta y)=(p/2,p/2)$, whose displacement magnitude is $p/\sqrt{2}$; it is not an equal-distance 45° comparison with the x and y curves. Left: selected-field infidelity. Right: selected power relative to the zero-registration reference. The same iris is used for the three directions at each beam size.')+r'''
The artificial no-carrier test recovers x/y symmetry to numerical precision. With a carrier, the difference is allowed and measured. Part of the x dependence is a throughput effect: at zero relative registration, the two blazes occupy coincident staircases; shifting one lattice divides their summed phase step between staggered boundaries. This redistributes diffraction efficiency even when the retained envelope remains nearly unchanged. Common-mode motion leaves the two staircases mutually aligned and differs from shifting only one panel.

Carrier orientation and sign on the physical panels are still unverified. These directional results apply to the declared x-carrier reference; rotation or parity in the actual inter-panel transfer can rotate or change the laboratory signature. They do not establish that an observed camera x-axis asymmetry comes from pixel registration.
''')
    page('6.4 Pixel corner edge and centre positions',r'''
The complete unit-cell maps in Figure 6 prevent an x-only sweep from being mistaken for a global tolerance. They contain 64 distinct positions $(\rho_x,\rho_y)\in\{0,1/8,\ldots,7/8\}^2$ for each selected beam size. The full-pitch endpoint is omitted because it repeats zero.
'''+figure('05_unit_cell','Vortex-owner registration maps for L = 20. The top row reports selected-field infidelity; the bottom reports selected power ratio. Each panel has its own labelled colour scale because the sensitivity differs markedly between radii. The maps are quantitative supplements to the large actual XY figures, not substitutes for them.')+rf'''
For 250 pixels, the largest selected-field infidelity in the sampled cell is {unit[unit.w_px==250].raw_infidelity.max():.2e}; its selected power ratio spans {unit[unit.w_px==250].power_ratio.min():.6f}–{unit[unit.w_px==250].power_ratio.max():.6f}. The corresponding maps for 50 pixels show larger deformation. Pixel-centre registration is not universally the best state: the optimum depends on the command and the chosen metric. The reference is a declared corner registration rather than an experimentally established optimum.

The unit-cell maps are periodic local-lattice statements. Real translations can additionally change finite-aperture clipping, pointing, illumination and the inter-panel transfer. Those effects require independent controls rather than being folded into the registration variable.
''')
    page('7.1 Stress beam immediately before the axicon',r'''
The following four figures use $L=20$ and $w=50p=0.4$ mm. The adjusted iris retains the intended vortex spectrum sufficiently to produce a dark-centred propagated ring. This is the representative registration stress case, rather than the near-total rejection imposed by the old fixed narrow iris.
'''+figure('stress_pre_intensity','Actual XY intensity immediately before the physical axicon after ideal-4F selection. Five x registrations share one physical crop and one colour scale fixed to the zero-shift peak. Neither panel intensity nor selected power is independently renormalised.')+r'''
The residual and complex-field measurements are essential because the raw envelope can look similar while the selected phase changes. Immediate intensity on a phase-only panel is not an appropriate standalone registration detector. Nor does an annular intensity pattern prove the requested topological charge. The retained complex fields are used to check winding on an identified contour in the propagated plane.
''')
    page('7.1 Signed stress residual before the axicon',r'''
Subtracting the zero-registration intensity exposes the spatial change that is difficult to judge from independently viewed profiles. Figure 8 uses a common symmetric residual scale. Positive and negative lobes distinguish local redistribution from an overall brightness gain.
'''+figure('stress_pre_residual',r'Signed pre-axicon residual $(I-I_0)/I_{0,\max}$ for the same five registrations, physical crop and reference as Figure 7. The reference panel is identically zero. The symmetric colour range is shared by all residuals.')+r'''
An absolute residual includes transmission change. The power-normalised residual reported separately compares $I/P$ with $I_0/P_0$ and is the preferred image-space indicator of shape change. Neither is a maximum-pixel percentage error: a low-intensity pixel can have a large relative error while contributing little optical power.

These panels describe the field after the selected channel has been reconstructed. The physical Fourier-plane iris itself contains a displaced focused spectrum, not this image-plane profile. Numeric selected coefficients are retained so that the Fourier plane and its transmitted spectrum can be regenerated without interpreting a rendered image as data.
''')
    page('7.2 Propagated stress intensity after the axicon',r'''
The physical axicon applies a conical phase, $t_{\mathrm{ax}}(r)=\exp(-ik_r r)$. Therefore $|E t_{\mathrm{ax}}|^2=|E|^2$ immediately behind an ideal phase-only surface. Figure 9 instead shows a finite propagation distance downstream, chosen from the zero-registration axial scan and held fixed across the comparison.
'''+figure('stress_post_intensity','Actual propagated XY intensity for L = 20, w = 50 pixels, at the labelled conditional-model plane. Five offsets use the same crop, reference and colour scale. The model assumes a sharp effective cone, not a measured apex or a full vector surface-refraction calculation.')+rf'''
The half-pixel x state changes the propagated peak by {pct(s20.peak_ratio-1)} at this plane and gives a {pct(s20.power_normalised_residual_l2)} power-normalised intensity residual. These are measurable deformation indicators, but they do not describe catastrophic destruction of this 50-pixel ring. The more severe 25-pixel case is examined separately. The intended winding is checked on the retained scalar complex field, not inferred from the dark core alone.
''')
    page('7.2 Signed propagated stress residual',r'''
The propagated residual in Figure 10 resolves weak azimuthal and radial changes. The axicon and propagation redistribute a local phase error into the ring system; registration does not merely translate an otherwise unchanged intensity distribution.
'''+figure('stress_post_residual',r'Signed propagated residual at the same plane as Figure 9. All panels share a symmetric colour scale and the zero-state normalisation. The growing structured residual indicates deformation of the selected vortex field, alongside its small throughput change.')+r'''
The scalar propagation uses the angular-spectrum factor

$$H(f_x,f_y;z)=\exp\!\left[iz\sqrt{k^2-(2\pi f_x)^2-(2\pi f_y)^2}\right].$$

The common carrier phase is omitted from numerical comparisons because it is a global phase. Propagating frequencies are retained; no physical loss is attributed to numerical filtering. The absence of an objective and material interface bounds the claim to free-space source morphology. Fine focal-component or ultrafast material-response conclusions are outside this study.
''')
    page('7.1 Bench size intensity before the axicon',r'''
The next four figures repeat the L = 20 sequence for $w=250p=2$ mm. Their role is to answer the current operating question directly. The adjusted iris is fixed for this beam size throughout its registration sweep. It differs from the stress-case iris because the intended vortex spectrum is narrower.
'''+figure('bench_pre_intensity','Actual XY intensity immediately before the axicon for the bench-size beam. The five offsets share the same physical crop and scale. Differences from the 50-pixel case are not concealed by independent panel normalisation.')+rf'''
The half-pixel x state has selected-field infidelity {b20.raw_infidelity:.2e} and selected power ratio {b20.power_ratio:.6f}. Small profile differences coexist with a roughly 2.5% increase in selected power. This is largely the two-blaze registration effect described in Section 6.3. It should not be reported as improvement of the vortex architecture: a phase allocation winner would require panel-specific physics absent from the model.
''')
    page('7.1 Bench size signed residual before the axicon',r'''
The bench-size residual is deliberately shown on a separate scale from the stress sequence so that its spatial signature remains visible. Its colourbar states the magnitude; a similar-looking red-blue pattern on a different scale does not imply equal sensitivity.
'''+figure('bench_pre_residual',r'Bench-size signed pre-axicon residual. The colour range is symmetric and fixed throughout this sequence. It exposes weak redistribution and the selected-channel brightness change without exaggerating them through independent panel rescaling.')+rf'''
After normalising by each field's retained power, the half-pixel x intensity residual before the axicon is {pct(b20.power_normalised_residual_l2)}. The raw residual is larger because it includes the selected-power gain. This distinction is needed when comparing camera images acquired with automatic exposure or plotted after per-frame normalisation. An exposure-normalised image can hide a real power effect even while correctly revealing shape stability.
''')
    page('7.2 Bench size propagated intensity',r'''
Figure 13 is the principal visual bench result. The same reference-derived downstream plane is used for all five states. The ring radius and dark core remain stable in the conditional cone calculation; the brightness change is small on the shared scale.
'''+figure('bench_post_intensity','Actual propagated XY intensity for L = 20 and the approximately 2 mm beam radius. All offsets use the labelled finite distance downstream of the physical axicon. The plot is not an immediate post-surface intensity image.')+rf'''
The half-pixel x peak changes by {pct(f20.peak_ratio-1)}, while the power-normalised image residual is {pct(f20.power_normalised_residual_l2)}. The intensity therefore remains robust despite a measurable local complex-field change at the panels. The source calculation supports relative registration sensitivity at this beam size; the plotted absolute ring dimensions and z coordinate remain conditional on the nominal relay and cone.
''')
    page('7.2 Bench size propagated signed residual',r'''
The residual in Figure 14 makes the remaining change inspectable. A radial oscillation pattern can reflect small ring redistribution even when the central feature appears unchanged. The residual should be read alongside its scale and the power-normalised metric.
'''+figure('bench_post_residual',r'Signed propagated bench-size residual relative to the zero-registration state at the plane in Figure 13. The same symmetric scale applies to all offsets. The reference image is zero by construction.')+r'''
Centroid recentering is reported but is not used to maximise overlap. A centroid shift is estimated from intensity and then applied to the **complex** field by a Fourier translation. This operation can make infidelity larger: it also translates a rapidly varying ring phase and does not remove phase tilt or deformation. There is no mathematical inequality requiring centroid-registered fidelity to exceed raw fidelity. A large centroid-registered value alone is therefore not evidence of a grossly deformed intensity profile. Raw field fidelity and image-space residuals remain the principal comparisons.
''')
    page('7.3 Axial consequences of registration',r'''
The axial scan evaluates the peak transverse intensity inside a fixed ±150 μm ROI, not the on-axis intensity of a vortex. Figure 15 compares zero and half-pixel registration at the same z samples. It distinguishes a small amplitude change from displacement or shortening of the axial envelope.
'''+figure('08_axial_evolution','ROI peak versus distance downstream of the conditional nominal axicon. Solid curves are zero registration; dashed curves are half-pixel x registration. Curves of each charge are normalised to their zero-registration axial maximum. Stress and bench radii occupy separate panels.')+table(['Bench charge','Zero-state peak z','Half-maximum sample interval','z step'],[
        [f'L = {L}',f'{zones[(zones.w_px==250)&(zones.L==L)&(zones.offset_px==0)].iloc[0].z_peak_m*1e3:.3f} mm',
         f'{zones[(zones.w_px==250)&(zones.L==L)&(zones.offset_px==0)].iloc[0].start_sample_m*1e3:.3f}–{zones[(zones.w_px==250)&(zones.L==L)&(zones.offset_px==0)].iloc[0].end_sample_m*1e3:.3f} mm',
         f'{zones[(zones.w_px==250)&(zones.L==L)&(zones.offset_px==0)].iloc[0].z_step_m*1e3:.3f} mm'] for L in [10,20]])+r'''
The half-maximum interval is a sampled envelope diagnostic, not proof of propagation-invariant Bessel-zone length. An unchanged interval only excludes changes large enough to cross these samples. Late stress-case samples can contain boundary return; the boundary-power ledger identifies those intervals and they must not be used to establish a quantitative zone. The headline plane is separately convergence-checked.
''')
    page('8 Severe mechanism stress case',r'''
The 25-pixel-radius beam compresses L = 20 onto a radius of only 0.2 mm. Figure 16 gives the presentation-scale comparison requested for a visibly sensitive beam. Both rows use an adjusted iris held fixed during that row's sweep; this is not the near-zero-transmission fixed-iris artefact.
'''+figure('presentation_04_severe_stress_vs_bench','Top: the deliberately severe 25-pixel-radius case at its labelled reference-derived plane. Bottom: the 250-pixel bench-size field. Zero, quarter- and half-pixel x offsets share one crop and reference scale within each row. Different baseline planes and iris bandwidths are used for the two radii; this compares registration sensitivity within each case, not absolute brightness between beams.')+rf'''
The severe case has a half-pixel x shape residual of {pct(stress.iloc[-1].power_normalised_residual_l2)} and a peak ratio of {stress.iloc[-1].peak_ratio:.3f}. Its nonuniform ring is already a sampling stress at zero registration; it must not be presented as the current laboratory beam. The 50-pixel case sits between this limit and the bench regime. Increasing the radius reduces registration-induced redistribution substantially.

The fixed-iris control at 50 pixels tells a different story: its zero-state selected fraction is only approximately $7.44\times10^{-5}$ and its propagated central field has winding zero on the inspected contours. Calling that field a degraded L = 20 vortex-Bessel beam would be misleading. Adjusting the iris separates this filtering failure from the registration deformation shown here.
''')
    page('9 Dimensionless interpretation and its limits',r'''
The azimuthal phase gradient suggests the local dimensionless parameter $\eta=Lp/w=L/w_{\mathrm{px}}$. It compares a typical phase variation across a pixel with the beam radius. The singular central region complicates a simple quadratic expansion because $L/r$ is unbounded at the origin. A finite Gaussian weighting and a finite lattice make an empirical exponent plausible, not universal.
'''+figure('06_scaling','Left: no-carrier local half-pixel x diagnostic, isolating the vortex sampling scale. Right: the primary two-blaze selected-field results with the adjusted iris. Colours indicate beam radius. The fitted local law is not applied to final propagated intensity or extrapolated outside the sampled matrix.')+rf'''
For the carrier-free local diagnostic over the sampled charges and radii, a log-space fit gives $1-F\simeq{fit['prefactor']:.3f}\eta^{{{fit['exponent']:.3f}}}$ with log-space $R^2={fit['log_R_squared']:.4f}$. This describes that diagnostic matrix only. The carrier introduces a phase-step floor; quantisation, square-lattice harmonics and the filter introduce additional dimensionless variables. The primary two-blaze filtered data do not follow a single universal η law.

A second useful parameter is $\chi=L/(2\pi Bw)$. It compares the vortex gradient at the beam radius with the filter bandwidth. The adjusted rule bounds χ near 0.2 for large enough L; the fixed 2.5 kcycles/m iris can instead place high-charge small beams in a strong-rejection regime. The local η scaling and the Bessel-forming/filtering condition must therefore be considered separately.
''')
    benchrows=[]
    for L in [10,20]:
        b=row(L);f=row(L,df=fine)
        benchrows.append([f'L = {L}',f'{b.raw_infidelity:.2e}',pct(f.peak_ratio-1),pct(f.power_normalised_residual_l2)])
    page('10 Current approximately 2 mm bench answer',r'''
For the declared phase-only effective-channel baseline, subpixel registration is a low-priority explanation for gross destruction of the current 2 mm-radius Bessel ring. Both L = 10 and L = 20 remain robust over the requested 0–half-pixel x sweep. Small field changes, selected-power changes and resolvable residual structure remain present.
'''+table(['Charge','Selected-field infidelity at half pixel','Propagated peak change','Propagated shape residual'],benchrows)+figure('10_bench_answer','The full 0–half-pixel sequence for the 250-pixel beam. Left: propagated peak change. Right: power-normalised intensity residual. Both are evaluated at the fixed reference-derived plane of each charge; the underlying fine complex fields are retained.',6.45)+r'''
These percentages are conditional simulation outputs, not measured tolerances. They do not excuse beam–hologram decentre, axicon misalignment, camera aliasing, LUT error or an unknown inter-SLM transfer. Those effects can be much larger and are not swept here. The practical conclusion is to establish geometry, camera sampling and phase calibration before attempting submicron panel positioning solely to repair a strongly distorted bench-size ring.

The fixed-iris and varied-bandwidth controls also keep the bench selected-field infidelity small. This supports the local robustness interpretation across the tested selection widths, while the actual hardware pinhole still needs measurement. The choice of an adjusted iris clarifies the mechanism; it does not make its physical diameter known.
''')
    page('11 Panel roles and architecture symmetry',r'''
The two panel operators in the flat-correction identity-transfer model are multiplicative. With identical panel response, exchanging their physical labels and their registration states leaves the product unchanged:

$$A:\ t_{v}(\rho_1)t_{b}(\rho_2),\qquad B:\ t_{b}(\rho_2)t_{v}(\rho_1).$$

Here $t_v$ includes vortex plus blaze and $t_b$ includes blaze plus zero correction. The selected field and every common downstream operation are consequently identical. This is stronger than an empirical B/A ratio near unity: the numerical implementation checks equality of the complex coefficients themselves.
'''+figure('09_panel_roles','L = 20 registration of the vortex owner, blaze owner, common mode and differential mode. Architecture A uses SLM1 as vortex owner; Architecture B uses SLM2. Equivalent physical roles match after panel exchange, rather than being compared under mismatched phase commands.')+rf'''
Across {len(roles)} paired role states, the maximum A/B coefficient difference is {roles.max_AB_coefficient_difference.max():.1e}. Thus A/SLM1-vortex matches B/SLM2-vortex, and A/SLM2-blaze matches B/SLM1-blaze. Common and differential modes are also included. At half-pixel offsets, differential and common configurations can become lattice-periodic equivalents; intermediate offsets distinguish them.

This equality is a validation result. It cannot establish that vortex generation on SLM2 is superior in the physical bench. A measured transfer $T_{{12}}$, panel-specific calibration, static phase or correction map can break the symmetry. A spatial correction cannot be assumed to behave like the flat blaze-only panel.
''')
    page('12 Correction maps and experimental discrimination',r'''
The calibration archive contains Beamage frames and response summaries for correction-gain and low-order-mode probes. The accompanying inference table describes coarse response directions, not a calibrated spatial phase command. The q20 retrieval bundles contain model-space correction predictions; their summaries explicitly state that the direct/conjugate branch, SLM2 coordinate mapping and 1030 nm phase LUT remain unresolved. The hardware calibration template has null scale, parity, rotation and conjugacy fields. No trustworthy panel-coordinate correction map was available for a correction-ON registration sweep.

This is not evidence that a real correction panel is insensitive. A nonflat correction has its own spatial gradients, which interact with its lattice registration. Correction-ON is the next model layer once the measured or fitted phase map has an accepted coordinate transform and provenance. The candidate retrieval images are not substituted for that missing calibration.

The laboratory test should first record both panel serials, operating wavelength, incidence and polarisation, driver/LUT ownership, static flattening maps, native mask hashes and blaze signs. Determine whether SLM1 is imaged onto SLM2 by observing asymmetric fiducials and measuring scale, rotation and parity. Record the intervening optics and separations. Inspect individual focal lengths and locate the Fourier plane; measure the iris radius and the selected-order centres in a calibrated frame. Confirm the axicon identity and measure its actual cone from a z scan rather than adopting the source-parity 2° optic.

Then perform the compensated translation on a stable mount while retaining $b=h$. A pattern-only fractional centre scan is a decentre experiment and is not an equivalent substitute. Acquire repeated, unsaturated camera frames at identical exposure and at matched pre-axicon and propagated planes. Record independent power so that the expected several-percent carrier-efficiency change is not mistaken for deformation. Use sufficient optical magnification or pixel integration to resolve the rings; detector fourfold artefacts must not be attributed to the SLM lattice without a camera-sampling control.

The discriminating signatures are a periodic response with one pitch, x/y differences tied to the carrier frame, a weak bench-size signed residual, and a stronger small-beam distortion. Interferometric or holographic measurement is needed to test complex fidelity directly; camera intensity alone tests the image-space metrics. Repeat the same role sweeps in both phase allocations after measuring the actual transfer. Only then can an architecture-ranking claim be made.
''')
    page('13 Limitations and conclusions',r'''
The reference is scalar, phase-only and linear. It includes exact lattice intersections, an ideal quantised command, a throughput-only fill factor, ideal source-frequency selection and an effective conical phase. It excludes measured fringing, LC angular response, unmodulated leakage, panel-specific LUTs, nonflat correction, a measured inter-panel transfer, lens aberrations, thick-surface vector refraction, apex defects, an objective, material response and a detector model. Wide diagnostic irises may compromise real diffraction-order isolation. Those omissions define the scope of the prediction rather than being treated as measured absences.

The 20° axicon description is interpreted **conditionally** as a base angle for the nominal cone calculation. Thorlabs gives the flat-entry cone-exit relation $\beta=\arcsin(n\sin\alpha)-\alpha$ [2]; $k_r=k\sin\beta$ then defines the effective transverse cone phase. The exact part number, orientation and index are not verified. The archived fitted $k_\perp=531200$ m⁻¹ gives a different cone and a roughly 42 μm L = 20 ring scale; it is itself model-inferred, not independent calibration. Absolute ring radius, z at peak and zone length in this report must not be relabelled as measured bench values.

Pure registration changes which portions of a continuous vortex and blaze are represented by each physical pixel. Higher charge and smaller illumination radius increase local sensitivity. The carrier gives a preferred direction and introduces efficiency changes that should be separated from deformation. With the iris adjusted, a 50-pixel L = 20 beam has a moderate visible change; a 25-pixel beam can exhibit substantial ring redistribution. The current 250-pixel beam remains robust in the tested baseline.

The axicon does not alter immediate intensity. It converts the retained phase difference into propagated structure; raw XY profiles and their signed residuals are therefore shown at finite downstream distance. The bench curves exhibit small paired intensity changes and no resolved change of the sampled axial half-maximum interval. Centroid recentering is a specified diagnostic, not guaranteed improvement of complex fidelity.

The present model cannot distinguish an intrinsic benefit of vortex-on-SLM1 from vortex-on-SLM2. Equivalent roles are exactly symmetric when both panels are blazed and panel-specific transfer is absent. The next decisive measurement is the inter-SLM transfer and coordinate transform, followed by calibrated phase response and a traceable correction map. Until then, the defensible conclusion concerns local pixel registration, not architecture superiority.
''')
    page('Appendix A Reproducibility and quantitative checks',r'''
The evidence package contains the primary adjusted-iris matrix, the fixed-iris control, full unit-cell maps, independently paired architecture-role states, axial and propagated metrics, topology diagnostics, numerical power ledgers, convergence controls and complex NPZ fields. Manifests record source hashes, numerical parameters, branch baseline and SHA-256 hashes of the numeric evidence. Images are generated from those arrays and CSVs; no reported metric is measured from a PNG.

The main reproduction route is:

```text
PYTHONPATH=. python tools/run_registration_reference_study.py
PYTHONPATH=. python tools/run_registration_extra_controls.py
PYTHONPATH=. python tools/complete_registration_diagnostics.py
PYTHONPATH=. python tools/plot_registration_reference_study.py
PYTHONPATH=. python tools/build_registration_reference_report.py
```

The fixed control uses `--iris-policy fixed --out outputs/validation/registration_definitive`. The full command sequence, dependencies and stage options appear in `docs/100_slm_registration_definitive.md`. Heavy NPZ, rendered figures and report files remain outside Git and are supplied as deliverable or workflow artifacts. Source, tests and lightweight provenance stay on the architecture-comparison branch.

The primary raw field fidelity is $|\langle E,E_0\rangle|^2/(\langle E,E\rangle\langle E_0,E_0\rangle)$. Full selected-field values use Fourier coefficients and Parseval's theorem. Post-axicon values use the explicitly declared ±150 μm complex-field ROI. The image residual is $\|I-I_0\|_2/\|I_0\|_2$; the shape residual applies the same expression to $I/P$ and $I_0/P_0$. Their percentages are L2 ratios, not maximum-pixel errors.

Ring radii, widths, core intensity, ellipticity, azimuthal purity and enclosed energy are retained as scalar diagnostics. Their physical scales remain conditional. The width is a sampled radial half-maximum interval, not a sub-grid fit. Enclosed energy is normalised to the retained ROI, not total propagated power. The winding contour and its radius are recorded; zero-amplitude contours require caution. These details prevent secondary metrics from exceeding their numerical or experimental maturity.
'''+table(['Convergence item','Check or evidence'],[
        ['Pixel boundaries','Analytic rectangle integrals; independent 4/8/16-sample raster test'],
        ['Illumination','q = 1, 2, 4; headline q = 2 with q = 4 control'],
        ['Cone sampling','Approximately 2 μm axial scan; 1.25 μm headline profiles'],
        ['Bench window','8 mm versus full 12 mm source-frequency window at L = 20'],
        ['Architecture exchange','Exact equality across paired complex coefficients']]))
    page('References and evidence provenance',r'''
[1] HOLOEYE Photonics, “PLUTO-2.1 LCOS Spatial Light Modulator.” Manufacturer specification for 1920 × 1080 addressing and 8 μm pitch. https://holoeye.com/products/spatial-light-modulators/pluto-2-1-lcos-phase-only-refl/ (accessed 1 October 2026).

[2] Thorlabs, “Axicons, UV Fused Silica.” Manufacturer base-angle and cone geometry; the page identifies 20° physical-angle axicons and the relation $\beta=\arcsin(n\sin\alpha)-\alpha$. https://www.thorlabs.com/newgrouppage9.cfm?objectgroup_id=4277 (accessed 1 October 2026). This reference establishes the convention used for the conditional calculation, not the identity of the user's optic.

[3] K. Matsushima and T. Shimobaba, “Band-limited angular spectrum method for numerical simulation of free-space propagation in far and near fields,” Optics Express 17, 19662–19673 (2009). https://doi.org/10.1364/OE.17.019662. The numerical propagation audit distinguishes bandwidth restriction from physical transmission.

[4] J. W. Goodman, Introduction to Fourier Optics, 3rd ed., Roberts and Company (2005). Fourier-transform lens and ideal coherent 4F filtering framework. The pixel-rectangle transform and fidelity expressions are derived explicitly in this report.

## Project evidence inspected

The starting branch was `slm-registration-architecture-comparison`, baseline commit `a946ddd5430427734d9d59ccb2d85745d76d247f`. The existing one-blaze predecessor was not used as a physical architecture comparison. Relevant source modules were `slm_pixel_registration.py`, `slm_registration_architectures.py`, `slm_registration_metrics.py`, `vortex_explicit_4f.py`, `vortex_system_route.py`, `slm_model.py` and the registration runners, figures, report builders and tests.

The geometry search inspected the canonical hardware manifest; `cslm_physical_axicon_bench_inventory.json`; the nominal F300 profile; the unfilled measured-bench template; and the bench evidence register. The active inter-panel distance in the demo inventory is 0.04 mm and is explicitly a diagnostic placeholder. It was not imported as laboratory geometry.

The uploaded `Publication_Study (2).zip` is a June source archive. The archived `slm_lab_gui_v0_6_measurement_sessions.zip` confirms direct panel-coordinate phase evaluation and calibration hooks. `Practical_Optical_System_Volume_II.pdf` explicitly distinguishes the 300 mm lens separation from individual focal lengths and the 20° physical-optic description from the historical 2° source model.

`Calibration.zip` contains response acquisitions, not a trustworthy panel-coordinate correction phase array. `current_calibration_inference.csv` records coarse correction-gain and modal response directions. The q20 local retrieval and simulation bundles describe uncalibrated model-space predictions; their template leaves conjugacy, scale, rotation, parity and LUT calibration unresolved. These materials support the stated missing-calibration boundary. They do not supply the missing transfer or justify correction-ON production results.

The short reproduction note and machine-readable evidence manifests accompany this report. The three principal presentation outputs are the severe-stress versus bench XY comparison, the bench ring with signed residual, and the bench intensity-change graphs. Their captions retain the same conditional-model and stress-case labels as the academic figures.
''')
    md=out/'SLM_Pixel_Registration_Report.md';md.write_text(BREAK.join(pages).replace('\\\\','\\'))
    docx=out/'SLM_Pixel_Registration_Report.docx'
    subprocess.run(['pandoc',str(md),'-f','markdown+raw_attribute','--standalone','-o',str(docx)],check=True)
    doc=Document(docx)
    sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.75);sec.left_margin=sec.right_margin=Inches(.8)
    for style in doc.styles:
        if style.type==1:
            style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
    stylemap={s.name:s for s in doc.styles}
    normal=stylemap['Normal'];normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.06;normal.paragraph_format.space_after=Pt(6)
    for body in ['Body Text','First Paragraph']:
        stylemap[body].base_style=normal;stylemap[body].font.size=Pt(10.5)
    for name,size in [('Heading 1',16),('Heading 2',12),('Title',23)]:
        stylemap[name].font.size=Pt(size);stylemap[name].font.color.rgb=RGBColor(0,0,0)
        stylemap[name].paragraph_format.space_before=Pt(0);stylemap[name].paragraph_format.space_after=Pt(10)
    doc.paragraphs[0].style=stylemap['Title']
    for para in doc.paragraphs:
        if para.text.startswith('Figure '):
            para.style=stylemap['Caption'];para.paragraph_format.space_after=Pt(8)
            for r in para.runs:r.font.size=Pt(9);r.font.italic=False
        if para._p.xpath('.//w:drawing'):
            para.paragraph_format.keep_with_next=True;para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if para.style.name=='Normal' and para.text and not para.text.startswith('['):
            para.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for t in doc.tables:
        t.autofit=True
        for i,r in enumerate(t.rows):
            for c in r.cells:
                pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
                pr.append(borders)
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED' if i==0 else 'FFFFFF');pr.append(sh)
                for para in c.paragraphs:
                    para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(4)
                    for run in para.runs:run.font.size=Pt(9);run.font.bold=i==0
        hdr=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(hdr)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fr=OxmlElement('w:r');ft=OxmlElement('w:t');ft.text='1';fr.append(ft);field.append(fr);footer._p.append(field)
    doc.core_properties.title='Subpixel registration in a dual SLM vortex Bessel system'
    doc.core_properties.author='Samuel McKenna';doc.core_properties.subject='Audited phase-only registration sensitivity'
    doc.save(docx)
    (out/'report_evidence_check.json').write_text(json.dumps(dict(planned_pages=len(pages),figures=fignum,
        primary_rows=len(sweep),role_pairs=len(roles),unit_cell_states=len(unit),
        headline_values={'L10':dict(pre_infidelity=float(b10.raw_infidelity),post_peak_ratio=float(f10.peak_ratio),post_shape=float(f10.power_normalised_residual_l2)),
                         'L20':dict(pre_infidelity=float(b20.raw_infidelity),post_peak_ratio=float(f20.peak_ratio),post_shape=float(f20.power_normalised_residual_l2))}),indent=2))
    print(docx)


if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('--data',type=Path,default=ROOT/'outputs/validation/registration_iris_adjusted')
    a.add_argument('--figures',type=Path,default=ROOT/'outputs/figures/registration_definitive')
    a.add_argument('--out',type=Path,default=ROOT/'outputs/reports/registration_definitive');args=a.parse_args();run(args.data,args.figures,args.out)
