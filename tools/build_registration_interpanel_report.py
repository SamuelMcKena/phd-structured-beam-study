#!/usr/bin/env python3
"""Evidence-derived academic revision with editable native Word equations."""
import argparse,json,subprocess
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from vbb_study.digital_twin.registration_reference import spectrum_fidelity

BREAK='\n\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n\n'

def run(data,out):
    out.mkdir(parents=True,exist_ok=True);figs=data/'figures'
    m=pd.read_csv(data/'metrics.csv');c=pd.read_csv(data/'convergence.csv');pc=pd.read_csv(data/'post_convergence.csv')
    extra=pd.read_csv(data/'additional_diagnostics.csv');power=pd.read_csv(data/'post_convergence_power.csv')
    headline=m[(m.role=='vortex + blaze')&(m.offset_px==.5)]
    fine=pc[pc.control=='dx1um'];pages=[];fignum=0
    def page(title,text):pages.append('# '+title+'\n\n'+text.strip()+'\n')
    def figure(name,caption,width=6.6):
        nonlocal fignum
        fignum+=1
        return f'\n\n![]({(figs/name).resolve()}){{width={width}in}}\n\nFigure {fignum}. {caption}\n\n'
    def table(headers,rows):
        return '\n|'+'|'.join(headers)+'|\n|'+'|'.join(['---']*len(headers))+'|\n'+''.join('|'+'|'.join(map(str,r))+'|\n' for r in rows)+'\n'
    def arch(a):return 'A' if a=='upstream_vortex' else 'B'
    def get(L,w,a,df=headline):return df[(df.L==L)&(df.w_px==w)&(df.architecture==a)].iloc[0]
    bench=fine[fine.w_px==250];stress=fine[fine.w_px==50]
    page('SLM pixel registration with free space between two panels',f'''
Samuel McKenna  
Applied Optics and Photonics Group, Heriot-Watt University  
Technical research report, October 2026

## Abstract

This report revises the architecture comparison following the clarification that the two SLMs are separated by approximately 200 mm of free space and that the 4F order-selection system follows both panels. Each SLM carries a 20-pixel blaze on an 8 µm lattice. Registration changes the lattice beneath a fixed beam and intended continuous hologram; no beam–hologram decentre is imposed. Correction is zero throughout: the FLAT-CORRECTION PHYSICS-ISOLATION BASELINE.

The 200 mm transfer breaks the exact phase-allocation symmetry of the previous identity-transfer reference. That earlier equality remains a valid local-model validation, but cannot be applied to separated panels. With aligned same-sign carriers, scalar angular-spectrum propagation, an adapted downstream iris and a conditional ideal axicon, a half-pixel x shift of the vortex panel gives propagated intensity-shape residuals of {100*stress.power_normalised_residual_l2.min():.2f}–{100*stress.power_normalised_residual_l2.max():.2f}% for the 50-pixel stress radius. For the 250-pixel bench radius and charges 10 and 20, the corresponding finely sampled residuals are {100*bench.power_normalised_residual_l2.min():.3f}–{100*bench.power_normalised_residual_l2.max():.3f}% at the specified diagnostic planes. These are relative scalar predictions, not measured bench tolerances.

## Principal numerical observations
'''+table(['Charge and radius','Architecture','Selected infidelity','Fine propagated shape'],[
        [f'L={int(r.L)}, w={int(r.w_px)} px',arch(r.architecture),f'{r.selected_infidelity:.4g}',
         f'{100*get(r.L,r.w_px,r.architecture,fine).power_normalised_residual_l2:.3f}%'] for r in headline.itertuples()])+'''
The final ring intensity is substantially less affected in the bench-radius cases than in the deliberately small-beam stress cases. No architecture winner is inferred from different nominal fields or different diagnostic planes. The experimental transfer coordinates, panel response and correction phase still require calibration. The following sections distinguish model validation, simulated observation and experimental inference.
''')
    page('1 Experimental geometry and registration',r'''
Architecture A places vortex plus blaze on SLM1 and flat correction plus blaze on SLM2. Architecture B reverses vortex ownership while retaining both blazes. Figure 1 shows the clarified order of components. There is no order-selection aperture between the panels. The approximate gap is user-reported rather than measured.
'''+figure('01_geometry_200mm.png','Both SLMs precede the 4F system. The arrows denote free-space transfer and downstream order selection; the physical axicon is after that selection.',6.3)+r'''
At each panel, beam centre $b$, intended continuous hologram centre $h$ and physical lattice origin $g$ define

$$d_{\mathrm{decentre}}=b-h,\qquad \rho=(h-g)/p\pmod 1.$$

Registration changes $g=-\rho p$ while the imposed nominal beam and hologram axes remain fixed. Compensated panel and pattern translations therefore cancel geometrically. A downstream centroid change caused by the upstream pixelated phase is recorded as a propagated outcome. It is not suppressed by recentering each state. An actual intensity centroid and a nominal optical axis need not coincide.

The nominal upstream carrier walks by 1.28628 mm over 200 mm. SLM2 is represented in one fixed coordinate frame on that nominal ray. The first carrier is demodulated only in this mathematical representation; both physical phase commands remain blazed.
''')
    page('2 Registration on the physical pixel lattice',r'''
The red marker in Figure 2 is the common beam and intended hologram centre. The physical pixel boundaries translate beneath it by $-\rho p$. The blue field marks the illuminated central region of a broad beam; the full Gaussian footprint extends outside this microscopic lattice view. No independent optical or hologram translation is introduced.
'''+figure('01_registration_geometry.png',r'The seven deterministic positions include x shifts of 0, 1/8, 1/4, 3/8 and 1/2 pixel, a half-pixel y shift, and the half-pixel corner-to-centre shift. The absolute corner/edge/centre labels follow the declared boundary-origin convention. The beam and intended phase origin remain coincident in every view.')+r'''
Changing only the electronic vortex centre relative to a stationary beam changes $b-h$ and measures decentre sensitivity. A registration measurement instead uses compensated coordinates so that the same centred continuous structure is sampled on a different physical lattice.

An upstream registration change can subsequently move or deform the propagated field incident on SLM2. That movement is a physical consequence of the altered pixel command, not an independently swept decentre parameter. Its centroid is recorded relative to the fixed nominal carrier-ray origin, including any static baseline offset. No per-state translation is applied to remove it from the optical route.

Whole-pitch periodicity is validated for the interior lattice in the stated modulo-one-pixel convention. Finite-panel edge motion is not silently treated as a new registration mechanism.
''')
    page('3 Numerical model and optical assumptions',r'''
The input field is $A(x,y)=\exp[-(x^2+y^2)/w^2]$, with $w$ the $1/e$ amplitude radius. Each panel commands a centre-sampled phase, quantised to 256 ideal levels, on a physical 8 µm pixel. Illumination is integrated with subcells, but the commanded phase is not refined into smaller fictitious pixels. Each panel contributes amplitude $\sqrt{0.93}$ as a throughput-only fill-factor approximation.

With carrier frequency $c=6250$ cycles/m, the upstream nominal displacement is

$$x_c=z\frac{\lambda c}{\sqrt{1-(\lambda c)^2}}.$$

In the fixed nominal-ray frame, the scalar transfer is

$$H(\nu_x,\nu_y)=\exp\{iz[k_z(\nu_x+c,\nu_y)-k_z(c,0)]+i2\pi\nu_x x_c\},$$

$$k_z(f_x,f_y)=\frac{2\pi}{\lambda}\sqrt{1-\lambda^2(f_x^2+f_y^2)}.$$

SLM1 pixel subcells have analytic top-hat Fourier factors. Their propagated field is evaluated on the actual shifted SLM2 coordinates using chunked Bluestein Fourier sums. SLM2 then applies its own pixelated phase and finite aperture. The downstream ideal 4F selects the second carrier in the demodulated envelope, corresponding physically to the sum of both carriers.

The diagnostic iris radius is $B=\max[2500,5L/(2\pi w)]$ cycles/m. It is fixed for all registration states and both architectures within each charge/radius comparison. Widening it exposes the intended stress-case spectrum, but real order purity and unmodulated leakage remain uncalibrated. The ideal relay has unit magnification with image parity unfolded. The 300 mm lens separation alone does not establish those properties.

The axicon is a conditional 20-degree BASE-angle ideal cone with index 1.458, not a measured prescription. A phase-only axicon leaves its immediate intensity unchanged. Every reported post-axicon intensity is propagated to a nonzero distance. The scalar solver supports broad relative morphology, not vector focal-detail or material-processing predictions.
''')
    page('4 Code audit and numerical validation',f'''
The prior one-blaze result compared unequal physical phase commands and could not support an architecture preference. Both current panel commands contain the 20-pixel blaze. The earlier sampled 4F route also lost approximately 9–55% of power in nominal free-space steps; those propagated conclusions remain rejected. The independent source-spectrum reference is retained for local registration and exact zero-transfer validation.

The new tests compare fractional-origin Fourier sums against direct sums, verify unit-modulus transfer, whole-pitch periodicity, fixed beam coordinates, both carrier roles and nominal walk-off. The artificial no-carrier rotation diagnostic uses a square aperture: the real rectangular aperture need not possess quarter-turn symmetry. The zero-distance branch uses the exact intersection-of-lattices reference, not a truncated numerical band pretending to be an identity operator. A finite-distance test demonstrates noncommutation of vortex multiplication and propagation.

The convergence controls expand the numerical source-frequency radius from 60000 to 80000 cycles/m and the padded width from 40 to 48 mm. Source and SLM2 illumination quadratures are independently increased from two to four subcells per axis. The numerical band is not an inserted optical iris. Its omitted total power is distinct from transfer drift and downstream aperture loss.
'''+table(['Case','Architecture','Primary infidelity','Range across controls'],[
        [f'L={L}, w={w}',arch(a),f'{g[g.control=="primary"].selected_infidelity.iloc[0]:.4g}',
         f'{g.selected_infidelity.min():.4g}–{g.selected_infidelity.max():.4g}']
        for (L,w,a),g in c.groupby(['L','w_px','architecture'])])+f'''
The maximum recorded downstream propagation power drift is {power.propagation_power_drift_fraction.max():.3g}, below the repository limit of 0.05. Conservation alone does not exclude periodic-window artifacts; the 12 mm propagation-window controls and the 1 µm sampling controls therefore remain separate checks. Small overlaps accumulate in double precision even when stored fields are complex64.

Pixel-maximum, ring-radius and azimuthal metrics are more sensitive to transverse sampling than integrated shape residuals. Their presence in a CSV is not evidence of convergence. Headline intensity-shape values use the 1 µm controls; unqualified claims about fine peak position or azimuthal structure are not made.
''')
    page('5 Local registration and dimensionless interpretation',r'''
The preserved local study sweeps $L=0,1,3,5,10,20$ and $w/p=25,50,100,150,250$, with x, y and diagonal offsets from zero to half a pixel and selected two-dimensional unit-cell maps. Its 450 sweep rows describe the two-blaze identity-transfer baseline. The local phase-overlap evidence remains applicable to the sampling mechanism; it is not silently relabelled as the separated-panel route.

The helical phase gradient is $|\nabla\phi_v|=L/r$. Across one pixel, its local phase increment is approximately $Lp/r$, motivating

$$\eta=Lp/w=L/(w/p).$$

Increasing charge at fixed radius increases the unresolved phase variation. Increasing beam radius reduces the fraction of power near the poorly sampled central region. The carrier supplies an additional preferred axis and a background registration response, so an unrestricted single power-law fit in $\eta$ is not a universal description of the two-blaze model. The earlier empirical fit is not reused as a law for the propagated architecture.

Free-space transfer introduces another dimensionless variable,

$$\chi=\lambda z/w^2.$$

At 200 mm, $\chi=1.286$ for the 50-pixel radius and $\chi=0.05145$ for the 250-pixel radius. The upstream stress field therefore evolves much more strongly before SLM2 even without registration. A collapse depending only on $\eta$ cannot generally describe both phase allocations. The finite aperture, carrier, iris and downstream cone introduce further scales.

In the boundary-origin convention, zero registration places the intended singularity at a lattice corner; half a pixel in one axis gives an edge and half a pixel in both axes gives a pixel centre. These labels specify the reference convention, not the measured laboratory registration. The deterministic geometry diagrams accompany the evidence bundle and show the grid moving beneath a fixed beam/hologram marker.

Complex-field fidelity is

$$F=\frac{|\langle E,E_{\mathrm{ref}}\rangle|^2}{\langle E,E\rangle\langle E_{\mathrm{ref}},E_{\mathrm{ref}}\rangle}.$$

Centroid translation is a prescribed correction, not maximisation of this overlap. It can remove a displacement while retaining phase tilt and deformation, and its infidelity can exceed the raw value. The power-normalised intensity residual separately measures visible shape change rather than throughput.

$$\epsilon_I=\left[\frac{\sum_\Omega(I/P-I_0/P_0)^2}{\sum_\Omega(I_0/P_0)^2}\right]^{1/2},\qquad P=\sum_\Omega I.$$

Here $\Omega$ is the explicitly declared transverse region. This L2 shape measure is not the largest per-pixel fractional difference, nor a loss of total optical power. Signed XY residuals instead use $(I-I_0)/I_{0,\max}$ and retain throughput changes.
''')
    local=data/'local_reference';fit=json.loads((local/'scaling_fit.json').read_text())
    page('6 Charge and radius in the local baseline',r'''
The complete local sweep is independently reproduced here with both blazes retained. Figure 3 compares unfiltered phase-product sensitivity against the selected complex field. The beam sizes are 25, 50, 100, 150 and 250 physical pixels in radius; the last corresponds to approximately 2 mm. The 25-pixel and 50-pixel cases deliberately expose sampling limitations rather than represent the current bench illumination.
'''+figure('03_local_charge_radius.png','Half-pixel x registration of the vortex-plus-blaze panel. Charges include the non-vortex L=0 carrier control. The two-blaze identity-transfer route isolates the local sampling and order-selection mechanism; it is not the 200 mm separated-panel prediction.')+r'''
The carrier-only control can change the unfiltered complex field even without a vortex. Within the ideal selected channel, a lattice displacement of a sampled periodic carrier can instead contribute mainly a uniform phase or throughput change. These effects explain why an unfiltered phase infidelity and a selected-channel intensity change need not have the same scale.

Higher charge increases the local helical phase gradient; a wider Gaussian reduces the power fraction in its poorly sampled central region. The response is not assumed strictly monotonic at every charge because quantisation, filtering and finite apertures can affect small residuals.
''')
    page('7 Testing the dimensionless collapse',f'''
Figure {fignum+1} re-evaluates the local scaling rather than transplanting the earlier empirical law into the corrected architecture. The artificial no-carrier diagnostic gives a fitted prefactor {fit['prefactor']:.3f}, exponent {fit['exponent']:.3f} and log-space coefficient of determination {fit['log_R_squared']:.4f} over the sampled matrix. It is an empirical description of those vortex-only half-pixel points.
'''+figure('04_local_scaling.png','Left: vortex-only local sampling with the carrier intentionally removed as a diagnostic. Right: the corrected two-blaze selected channel. The latter does not collapse onto the same universal curve; beam radius remains distinguishable at comparable eta.')+r'''
The iris bandwidth adds a filtering scale $L/(2\pi B w)$, while the separated-panel route adds $\lambda z/w^2$. A fit to vortex-only local infidelity therefore cannot predict the final Bessel intensity or rank phase allocation across the free-space gap. The scaling variable remains useful for identifying high-charge small-beam stress, not for replacing the optical calculation.
''')
    page('8 Pixel corners edges and centres',r'''
The two-dimensional unit-cell scan in Figure 5 resolves eight positions per axis for the selected local cases. Its colours show a scalar fidelity metric, not a beam-intensity image. The principal visual beam evidence remains the large XY profiles in the following sections.
'''+figure('05_local_unit_cell.png','L=20 local unit-cell maps for the 50-pixel stress beam and 250-pixel bench radius. Each map uses its own explicitly labelled infidelity scale. The physical lattice origin at zero is a corner in this convention; half a pixel in both axes places the singularity at a pixel centre.')+r'''
Registration cannot be reduced to a single distance from an arbitrarily named pixel centre when a carrier selects one axis. Corner, vertical-edge and horizontal-edge states sample different phase partitions. The additional separated-panel direction controls retain these distinctions without identifying the local unit-cell maps with the propagated 200 mm route.
''')
    page('Field incident on the second physical panel',r'''
Before downstream order selection, the upstream stress vortex has propagated 200 mm and encountered the independent SLM2 lattice. The images below are fields at that actual plane, rather than images inferred from a selected-channel radial profile. The fixed frame follows the nominal carrier ray and does not follow state-dependent centroid motion.
'''+figure('07_actual_slm2_plane.png','Architecture A stress case at zero and half-pixel SLM1 registration. Left: intensity immediately incident on SLM2. Right: wrapped phase immediately after SLM2. The first carrier is represented in a demodulated coordinate frame; the second panel still applies its own physical blaze. Darkness does not by itself establish a topological winding.')+r'''
In the ideal phase-only model, the instantaneous SLM2 modulation changes phase rather than redistributing intensity; the throughput factor uniformly scales it. Spatial intensity changes occur in propagation and order selection. The same principle holds at the physical axicon, whose immediate intensity is not presented as a new downstream result.
''')
    for w,a,planes in [(50,'upstream_vortex',['pre','post']), (250,'upstream_vortex',['pre','post'])]:
        panel=1;label='Stress test' if w==50 else 'Bench radius';r=get(20,w,a);fr=get(20,w,a,fine)
        for plane in planes:
            for residual in [False,True]:
                name=f'fields_L20_w{w}_{a}_panel{panel}_{plane}_'+('residual' if residual else 'intensity')+'.png'
                where='before the physical axicon' if plane=='pre' else f'{1000*r.z_m:g} mm downstream of the physical axicon'
                page(f'{label} L20 Architecture A '+('signed residual ' if residual else 'intensity ')+('before axicon' if plane=='pre' else 'after propagation'),f'''
The {w}-pixel incident radius corresponds to {w*8e-3:g} mm on SLM1. Only its vortex-plus-blaze panel is shifted in this sequence. Architecture A's second panel retains its own blaze with zero correction. Figure {fignum+1} shows the actual XY {'signed intensity difference' if residual else 'intensity'} {where} for all five requested x registrations.
'''+figure(name,('Signed residuals are divided by the zero-registration peak and share a symmetric colour scale. ' if residual else 'All five intensities are divided by one zero-registration peak; they are not individually peak-normalised. ')+
                'The physical crop is fixed. Display interpolation improves readability only; numerical metrics use the saved complex arrays. Both panels remain blazed.')+f'''
At half a pixel, the selected full-field infidelity is {r.selected_infidelity:.5g}. The pre-axicon power-normalised intensity residual is {100*r.pre_power_normalised_residual_l2:.3f}%. At the indicated downstream plane, the finely sampled residual is {100*fr.power_normalised_residual_l2:.3f}%. That downstream metric is evaluated in a fixed 300 µm-wide XY region, not over the entire transverse field.

{'This radius is deliberately non-representative of the current 2 mm beam. The displayed deformation tests the sampling mechanism and must not be described as bench degradation.' if w==50 else 'This is the current approximately 2 mm radius. The signed residual exposes small changes that are difficult to distinguish in the raw ring intensity. It does not imply that an experimentally distorted ring has been explained by registration.'}

The propagated plane is fixed throughout the sweep. It is a diagnostic plane, not a claimed measured Bessel-zone maximum. A dark centre alone is not proof of winding; scalar complex-field contour diagnostics are retained separately.
''')
    for w in [50,250]:
        r=get(20,w,'downstream_vortex');fr=get(20,w,'downstream_vortex',fine)
        page(('Stress test' if w==50 else 'Bench radius')+' L20 Architecture B propagated intensity',f'''
Here the vortex is generated locally on SLM2 after the approximate 200 mm gap. SLM1 has flat correction plus its own blaze. The shifted panel is SLM2; its incident field and intended continuous hologram axis remain fixed while the lattice moves. Figure {fignum+1} uses the same five registrations but a clearly labelled Architecture B diagnostic plane.
'''+figure(f'fields_L20_w{w}_downstream_vortex_panel2_post_intensity.png',f'Actual XY intensity propagated {1000*r.z_m:g} mm beyond the conditional axicon. The crop and colour scale are shared across this sequence; they are not a common absolute-intensity scale with Architecture A.')+f'''
The half-pixel selected-field infidelity is {r.selected_infidelity:.5g}, with finely sampled downstream shape residual {100*fr.power_normalised_residual_l2:.3f}%. This result does not establish that Architecture B is experimentally preferable. The nominal field, working-distance distribution and diagnostic plane differ when vortex ownership changes across a free-space gap.

Comparing registration tolerance requires each perturbed field to be compared against its own zero-registration route. Comparing A and B at a common physical plane is also provided by the axial diagnostics. Differences in nominal morphology or throughput are not by themselves improvements in registration tolerance.
''')
    page('Current approximately 2 mm bench answer',f'''
For the stated conditional geometry, the current 250-pixel radius remains substantially less registration-sensitive than the 50-pixel stress case. Figure {fignum+1} separates selected complex-field change from propagated intensity change over the full x sweep for charges 10 and 20. The charge dependence is visible in the field-level metric even where the final ring change is small.
'''+figure('presentation_02_bench_sensitivity_200mm.png','Field infidelity and image-shape residual for the bench radius. Each legend declares the physical diagnostic distance. The right-hand graph uses primary 2 µm fields; the table below gives independent 1 µm endpoint controls.')+table(['Charge','Architecture','z beyond axicon','Fine shape residual'],[
        [int(r.L),arch(r.architecture),f'{1000*r.z_m:g} mm',f'{100*r.power_normalised_residual_l2:.3f}%'] for r in bench.itertuples()])+'''
The simulated evidence does not make subpixel registration a compelling primary explanation for gross distortion of the present bench beam. The complex field changes slightly while the final sampled ring intensity is relatively robust. This conclusion is limited to flat correction, ideal panel response and the assumed optical mapping. It does not cover decentre, LUT error, camera aliasing, axicon tilt or a non-flat measured correction map.
''')
    direction=extra[extra.diagnostic.isin(['y','diagonal','common','differential'])]
    page('Directions common and differential registration',r'''
The x blaze breaks rotational invariance. Artificial no-carrier square-aperture symmetry is recovered in tests, while carrier-induced anisotropy is permitted in the physical baseline. The endpoint results below retain y and diagonal shifts, as well as common and differential motion. A diagonal half-pixel shift moves each coordinate by half a pixel and is not the same Euclidean displacement as an x-only shift.
'''+table(['Charge radius','Architecture','Mode','Selected infidelity'],[
        [f'L={int(r.L)}, w={int(r.w_px)}',arch(r.architecture),r.diagnostic,f'{r.selected_infidelity:.4g}']
        for r in direction[direction.L==20].itertuples()])+r'''
Axial comparisons are saved from 1 to 25 mm downstream of the conditional cone, with baseline and half-pixel fields at common z values. The interval is diagnostic, not a calibrated experimental Bessel-zone length. A threshold-defined high-intensity interval depends on the chosen region and may be censored by the sampled endpoints; no exact zone start, end or length is inferred from a sparse uncalibrated scan.

The axicon does not erase complex-field differences. Propagation redistributes their visible intensity effect, so a weak effect at one plane does not prove identical fields everywhere. Conversely, a field-level infidelity is not a percentage loss of ring intensity. CSVs retain centroid displacement, power ratio and absolute and normalised residuals to distinguish these statements.
''')
    page('Axial evolution at common physical planes',f'''
The axial calculation compares each perturbed field with its own nominal route at the same downstream distance. Figure {fignum+1} avoids comparing the architectures only at their separate chosen diagnostic planes. The common 300 µm-wide transverse region is held fixed. Its metric is conditional on the ideal scalar cone and may omit power outside that region.
'''+figure('06_axial_registration.png','Half-pixel x registration of the vortex owner, evaluated at common distances from 1 to 25 mm beyond the physical axicon. The broad axial trend tests whether a single-plane observation hides a much larger downstream change. The sparse scan does not establish a measured Bessel-zone start or end.')+r'''
Numerical window controls accompany the fixed-plane headline cases. Fine peak location, fitted width and winding near low-intensity contours need stronger sampling and a calibrated cone before they support an experimental axial-zone claim. A dark central pixel or a ring-shaped image is not treated as sufficient proof that a well-defined charge persists across every contour.
''')
    cross=[]
    for L,w in [(20,50),(10,250),(20,250)]:
        aa=np.load(data/f'fields_L{L}_w{w}_upstream_vortex_panel1.npz')['coefficients'][0]
        bb=np.load(data/f'fields_L{L}_w{w}_downstream_vortex_panel2.npz')['coefficients'][0]
        cross.append([L,w,f'{1-spectrum_fidelity(aa,bb):.5g}'])
    page('Architecture interpretation and correction',r'''
For identity transfer and equal flat-correction responses, phase multiplication commutes. Swapping physical roles therefore gives exactly equal A/B fields; this is explicitly tested. With 200 mm propagation, the relevant operators satisfy

$$\mathcal{P}_{z}\mathcal{M}_{v}\ne\mathcal{M}_{v}\mathcal{P}_{z}.$$

The nominal selected-field differences below demonstrate that the separated-panel calculation is no longer a panel relabelling. They do not measure registration tolerance and do not rank the laboratory architectures.
'''+table(['Charge','Radius px','Nominal A versus B infidelity'],cross)+r'''
The original motivation that a locally generated downstream vortex might be less exposed to a second registered panel remains a testable hypothesis, not an assumed result. Under the present assumptions, both vortex-owner cases retain strong bench-radius robustness. An apparent advantage at different z planes, through different baseline intensity profiles or through an omitted blaze would not answer the same physical question.

No accepted panel-coordinate measured correction map or calibrated panel-specific LUT is introduced. A flat correction leaves its panel effectively blaze-only. Low sensitivity there is a property of this isolation baseline; the real corrected panel may have appreciable phase gradients and a different registration response. Correction-ON with a measured map is the next layer, not a completed experimental validation.

The approximate gap enables a conditional architecture calculation, but the remaining coordinate rotation, parity, carrier orientations, relay magnification and panel responses determine whether that conditional route describes the bench. An architecture winner should be reported only after the competing routes share accepted physical calibration and a specified operating-plane criterion.
''')
    page('Laboratory test limitations and conclusions',r'''
A laboratory registration test should translate the panel relative to a fixed centred beam and compensate the intended hologram coordinates, keeping beam–hologram decentre fixed. Acquisition should record both panel allocations, selected power, the actual iris and camera scale, and XY profiles before the axicon and at repeatable downstream distances. A 0–4 µm scan in 1 µm steps implements the requested eighth-pixel sequence. Axial scans and repeated exposures can distinguish weak shape changes from drift or camera sampling artifacts.

The highest-priority geometry measurement is the optical path and nominal-axis mapping from SLM1 to SLM2, including distance, carrier sign/orientation and transverse coordinate relation. The downstream magnification, image-plane position, iris and axicon effective cone must then be recorded. Panel-specific 1030 nm phase response and a measured correction phase are required before correction-ON can be interpreted quantitatively.

The local mechanism is rapid sampled helical phase variation, concentrated near the singularity and increasingly important at high charge or small beam radius. The separated-panel system adds diffraction evolution and carrier walk-off; a single scaling variable cannot describe it universally. Both blazes are necessary in either architecture. Pure registration remains distinct from imposed decentre, while propagated centroid motion is an observable consequence.

For the declared approximately 200 mm conditional route, the current 2 mm-radius beam is robust in the tested half-pixel sweeps compared with the intentionally small-radius stress beam. This is a relative scalar conclusion. Absolute Bessel dimensions, detailed vector peak structure, real diffraction efficiency, corrected-panel sensitivity and material response remain outside the accepted evidence.

## Reproducibility appendix

The source is on `slm-registration-architecture-comparison`. The reproduction note is `docs/101_slm_registration_interpanel_propagation.md`. The runner stages are `local`, `sweep`, `controls`, `postcontrols` and `extra`. Numeric evidence includes CSV tables, complex-field NPZ arrays and manifests with hardware assumptions, source commit and SHA256 source hashes. Plotting uses these arrays directly; no image is used as a metric input. Large evidence remains outside Git and is retained as an artifact.

The earlier local report and evidence archive remain an identity-transfer reference. This revision changes their architecture scope, not their historical provenance. The two presentation figures provide an actual XY stress/bench comparison and the bench sensitivity graphs. Standalone pre-axicon and propagated intensity and signed-residual sequences remain separate readable figures.
''')
    page('Methods references and evidence provenance',r'''
The numerical propagation uses exact scalar angular-spectrum phase within a finite computational band. Matsushima and Shimobaba discuss the sampling errors that can occur in discrete angular-spectrum propagation and motivate checking computational support independently of energy conservation [1]. The present code additionally tests fractional-origin Fourier evaluation against direct sums and changes source support, quadrature and downstream sampling independently.

The fractional-coordinate evaluation uses Bluestein chirp convolution, an algorithm also documented for the chirp-z transform [2]. The implementation is validated directly; naming an algorithm is not substituted for its numerical convergence check. Physical pixel commands remain fixed at 8 µm even when illumination quadrature is refined.

[1] K. Matsushima and T. Shimobaba, “Band-limited angular spectrum method for numerical simulation of free-space propagation in far and near fields,” Optics Express 17, 19662–19673 (2009). DOI 10.1364/OE.17.019662.

[2] SciPy documentation, `scipy.signal.CZT`, version 1.15.0, notes on Bluestein evaluation and numerical accuracy. https://docs.scipy.org/doc/scipy-1.15.0/reference/generated/scipy.signal.CZT.html

The approximate 200 mm gap and downstream relay placement originate in the user's subsequent setup clarification. Pixel pitch, blaze period, wavelength and beam radius are declared hardware-scale inputs. Equal ideal LUTs, aligned coordinates, carrier signs, relay magnification and the cone prescription remain numerical assumptions. None is presented as a measured calibration.

The measured-map audit found no accepted panel-specific correction phase or coordinate transfer suitable for this calculation. The repository's GUI and digital-twin phase definitions informed the centred pixel command; a template or synthetic phase file was not promoted into experimental evidence. Historical reports and the earlier branch remain preserved for traceability.
''')
    md=out/'SLM_Registration_200mm_Report.md';md.write_text(BREAK.join(pages))
    docx=out/'SLM_Registration_200mm_Report.docx'
    subprocess.run(['pandoc',str(md),'-f','markdown+raw_attribute','--standalone','-o',str(docx)],check=True)
    doc=Document(docx);sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
    for s in doc.styles:
        if s.type==1:s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
    for name in ['Normal','Body Text','First Paragraph']:
        doc.styles[name].font.size=Pt(11);doc.styles[name].paragraph_format.line_spacing=1.04
        doc.styles[name].paragraph_format.space_after=Pt(7)
    for name,size in [('Title',22),('Heading 1',15),('Heading 2',12)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].paragraph_format.space_before=Pt(0)
    doc.paragraphs[0].style=doc.styles['Title']
    for p in doc.paragraphs:
        if p.text.startswith('Figure '):
            p.style=doc.styles['Caption']
            for r in p.runs:r.font.size=Pt(9.5);r.font.italic=False
        if p._p.xpath('.//w:drawing'):p.paragraph_format.keep_with_next=True
    for t in doc.tables:
        for i,row in enumerate(t.rows):
            for cell in row.cells:
                pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
                pr.append(b)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
                    for r in p.runs:r.font.size=Pt(9.5);r.font.bold=i==0
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');sec.footer.paragraphs[0]._p.append(field)
    doc.core_properties.author='Samuel McKenna';doc.core_properties.title='SLM pixel registration with free space between two panels'
    doc.save(docx)
    (out/'report_evidence_check.json').write_text(json.dumps(dict(planned_pages=len(pages),figures=fignum,headline_rows=len(headline),
        shape_range_bench=[float(bench.power_normalised_residual_l2.min()),float(bench.power_normalised_residual_l2.max())]),indent=2))
    print(docx)

if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('--data',type=Path,required=True);a.add_argument('--output',type=Path,required=True)
    q=a.parse_args();run(q.data,q.output)
